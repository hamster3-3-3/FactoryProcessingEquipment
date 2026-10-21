#!/usr/bin/env python3
"""Read the five detail sheets; preserve each source row without deduplication."""
import argparse, collections, json, hashlib
from datetime import datetime
from zoneinfo import ZoneInfo
from pathlib import Path
import openpyxl

EQUIPMENT = ['炊飯','煮麵','烤箱','炙燒','蒸箱','油炸機','二重釜','平釜(日)','平釜(台)','蛋皮機','壓力鍋','無使用調理設備','萬能炒食機']
SHEETS = ['18℃','4℃飯','4℃麵','4℃其他','冷凍品']
def clean(value):
    return str(value).strip() if value is not None else ''
def convert(folder, output):
    records, sources = [], []
    for factory in ['晉欣','新豐','大溪','長家']:
        candidates = [p for p in Path(folder).glob('*.xlsx') if not p.name.startswith('~$')]
        matches = [p for p in candidates if factory in p.stem]
        if len(matches) != 1:
            raise ValueError(f'{factory}: 請只放一份來源檔，目前 {len(matches)} 份')
        path = matches[0]
        book = openpyxl.load_workbook(path, data_only=True)
        source = {'factory': factory, 'file': path.name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'sheets': {}}
        for title in SHEETS:
            sheet = book[title]
            headers = {clean(c.value).replace('\n','').replace('（','(').replace('）',')'): c.column for c in sheet[3] if c.value is not None}
            required = ['半成品名稱','商品名稱','單位用量','製程率(填數字)','風險說明']
            missing = [name for name in required if name not in headers]
            if missing: raise ValueError(f'{path.name}／{title} 缺少必要欄位: {missing}')
            def cell(row, key):
                return sheet.cell(row, headers[key]) if key in headers else None
            def value(row, key):
                c = cell(row, key)
                return c.value if c else None
            count = 0
            for row in range(5, sheet.max_row+1):
                half = clean(value(row, '半成品名稱'))
                if not half: continue
                raw_status = clean(value(row, '狀態'))
                status = {'架上':'上架','上架':'上架','待上架':'待上架','下架':'下架'}.get(raw_status, '未標示')
                yc = cell(row, '製程率(填數字)')
                rate = yc.value if yc else None
                fmt = yc.number_format if yc else ''
                if isinstance(rate, (float,int)):
                    # Excel unescaped percent scales by 100; literal \\% does not.
                    rate = round(rate * (100 if '%' in fmt.replace('\\%', '') else 1), 6)
                else:
                    rate = clean(rate) or None
                eq = [name for name in EQUIPMENT if clean(value(row,name)).upper() in ['V','Ｖ','✓','✔','√']]
                amount = value(row,'單位用量')
                unit = sheet.cell(row, headers['單位用量']+1).value
                if amount is not None and not isinstance(amount,(float,int,str)): amount = str(amount)
                record = dict(id=f'{factory}-{title}-{row}',factory=factory,temp='18℃' if title=='18℃' else '冷凍' if title=='冷凍品' else '4℃',subtype=title[2:] if title.startswith('4℃') else '',status=status,rawStatus=raw_status,category=clean(value(row,'商品分類(細分品類)')) or '未分類',product=clean(value(row,'商品名稱')),half=half,amount=amount,unit=clean(unit),yieldRate=rate,rawYield=yc.value if yc else None,yieldFormat=fmt,equipment=eq,check=clean(value(row,'檢核結果')),risk=clean(value(row,'風險說明')),note=clean(value(row,'備註(指示書編號/版本)')),source=path.name,sheet=title,row=row)
                if not record['product']: raise ValueError(f'商品名稱缺漏: {record["id"]}')
                records.append(record); count += 1
            source['sheets'][title] = count
        sources.append(source)
    if not records: raise ValueError('四廠皆無有效明細，停止發布，請檢查Excel。')
    data = dict(title='商品半成品與調理設備資料庫',version='Excel 自動更新版',importedAt=datetime.now(ZoneInfo('Asia/Taipei')).date().isoformat(),equipment=EQUIPMENT,sources=sources,records=records)
    out = Path(output); out.mkdir(parents=True, exist_ok=True)
    serialized = json.dumps(data, ensure_ascii=False, separators=(',',':'))
    (out/'data.json').write_text(serialized,encoding='utf-8')
    (out/'data.js').write_text('window.DATABASE='+serialized+';\n',encoding='utf-8')
    print(json.dumps({'total':len(records),'factory':dict(collections.Counter(r['factory'] for r in records)),'status':dict(collections.Counter(r['status'] for r in records))},ensure_ascii=False))
if __name__ == '__main__':
    parser=argparse.ArgumentParser();parser.add_argument('folder');parser.add_argument('--output',default='.')
    args=parser.parse_args(); convert(args.folder,args.output)
